from __future__ import annotations

import logging
import re
from pathlib import Path
from typing import Callable

from .models import Document

logging.getLogger("pypdf").setLevel(logging.ERROR)

_LOGGER = logging.getLogger(__name__)

SUPPORTED_SUFFIXES = {".txt", ".md", ".markdown", ".rst", ".pdf", ".docx"}

DEFAULT_IGNORED_DIRS = {
    ".git",
    ".hg",
    ".svn",
    ".venv",
    "venv",
    "env",
    "node_modules",
    "__pycache__",
    "site-packages",
    ".cache",
    ".mypy_cache",
    ".pytest_cache",
}

_SENTENCE_RE = re.compile(r"(?<=[.!?;:])\s+")
_PARAGRAPH_RE = re.compile(r"\n\s*\n")


def discover_files(path: str | Path, ignore_dirs: set[str] | None = None) -> list[Path]:
    p = Path(path)
    if p.is_file():
        return [p] if p.suffix.lower() in SUPPORTED_SUFFIXES else []
    ignored = DEFAULT_IGNORED_DIRS if ignore_dirs is None else set(ignore_dirs)
    found: list[Path] = []
    for candidate in p.rglob("*"):
        if not candidate.is_file() or candidate.suffix.lower() not in SUPPORTED_SUFFIXES:
            continue
        parts = candidate.relative_to(p).parts[:-1]
        if any(part in ignored or part.startswith(".") for part in parts):
            continue
        found.append(candidate)
    return sorted(found)


def read_file(path: str | Path) -> str:
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix == ".pdf":
        return _read_pdf(p)
    if suffix == ".docx":
        return _read_docx(p)
    return p.read_text(encoding="utf-8", errors="replace")


def _read_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("PDF support requires pypdf: pip install pypdf") from exc
    reader = PdfReader(str(path))
    pages = []
    for page in reader.pages:
        pages.append(page.extract_text() or "")
    return "\n\n".join(pages)


def _read_docx(path: Path) -> str:
    try:
        import docx
    except ImportError as exc:
        raise RuntimeError("DOCX support requires python-docx: pip install python-docx") from exc
    document = docx.Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def _split_long(segment: str, size: int, overlap: int) -> list[str]:
    pieces = []
    start = 0
    step = max(1, size - overlap)
    while start < len(segment):
        pieces.append(segment[start : start + size])
        start += step
    return pieces


def _segments(text: str, size: int, overlap: int) -> list[str]:
    segments: list[str] = []
    for paragraph in _PARAGRAPH_RE.split(text):
        paragraph = paragraph.strip()
        if not paragraph:
            continue
        for sentence in _SENTENCE_RE.split(paragraph):
            sentence = sentence.strip()
            if not sentence:
                continue
            if len(sentence) > size:
                segments.extend(_split_long(sentence, size, overlap))
            else:
                segments.append(sentence)
    return segments


def chunk_text(text: str, size: int = 900, overlap: int = 150) -> list[str]:
    if size <= 0:
        raise ValueError("size must be positive")
    if overlap < 0 or overlap >= size:
        raise ValueError("overlap must be within [0, size)")
    text = text.replace("\r\n", "\n").strip()
    if not text:
        return []
    if len(text) <= size:
        return [text]

    segments = _segments(text, size, overlap)
    chunks: list[str] = []
    current: list[str] = []
    current_len = 0

    for segment in segments:
        if current and current_len + len(segment) + 1 > size:
            chunks.append(" ".join(current))
            tail: list[str] = []
            tail_len = 0
            for prev in reversed(current):
                if tail_len + len(prev) + 1 > overlap:
                    break
                tail.insert(0, prev)
                tail_len += len(prev) + 1
            current = tail
            current_len = tail_len
        current.append(segment)
        current_len += len(segment) + 1

    if current:
        chunks.append(" ".join(current))
    return [c.strip() for c in chunks if c.strip()]


def load_documents(
    path: str | Path,
    chunk_size: int = 900,
    chunk_overlap: int = 150,
    ignore_dirs: set[str] | None = None,
    on_progress: Callable[[int, int], None] | None = None,
) -> list[Document]:
    documents: list[Document] = []
    files = discover_files(path, ignore_dirs=ignore_dirs)
    total = len(files)
    for position, file_path in enumerate(files, start=1):
        try:
            text = read_file(file_path)
        except Exception as exc:
            _LOGGER.warning("Skipping %s: %s", file_path, exc)
            if on_progress:
                on_progress(position, total)
            continue
        for index, chunk in enumerate(chunk_text(text, chunk_size, chunk_overlap)):
            documents.append(
                Document(
                    id=f"{file_path.stem}-{index:04d}",
                    text=chunk,
                    source=str(file_path),
                    metadata={
                        "chunk_index": index,
                        "suffix": file_path.suffix.lower(),
                    },
                )
            )
        if on_progress:
            on_progress(position, total)
    return documents
